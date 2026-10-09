import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_horizontal_border(paragraph, color_hex="1B4D3E", sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), sz)
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)

def build_docx(md_path, docx_path):
    doc = Document()

    # Set page margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base colors
    COLOR_PRIMARY = RGBColor(11, 59, 36)      # #0B3B24 Forest Pine
    COLOR_SECONDARY = RGBColor(224, 133, 56)  # #E08538 Warm Amber
    COLOR_DARK = RGBColor(30, 41, 59)         # #1E293B Slate Dark
    COLOR_MUTED = RGBColor(100, 116, 139)     # #64748B Slate Muted

    # Setup styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = COLOR_DARK

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # Read frontmatter / title info
    # We will build an elegant Cover / Title Header
    # Page 1: Title Header matching the reference PDF exactly
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(120)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Stackly Travel & Tourism Website")
    run_title.font.name = 'Segoe UI Semibold'
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(4)
    p_sub.paragraph_format.space_after = Pt(24)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Technical Documentation")
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(16)
    run_sub.font.color.rgb = COLOR_SECONDARY

    p_tag = doc.add_paragraph()
    p_tag.paragraph_format.space_before = Pt(0)
    p_tag.paragraph_format.space_after = Pt(6)
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_tag = p_tag.add_run("Responsive HTML5 Template for Luxury Expeditions, Tour Operators & Travel Agency Portals")
    run_tag.font.name = 'Segoe UI'
    run_tag.font.size = Pt(11)
    run_tag.font.italic = True
    run_tag.font.color.rgb = COLOR_MUTED

    p_ver = doc.add_paragraph()
    p_ver.paragraph_format.space_before = Pt(4)
    p_ver.paragraph_format.space_after = Pt(180)
    p_ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_ver = p_ver.add_run("Version 1.0")
    run_ver.font.name = 'Segoe UI'
    run_ver.font.size = Pt(10.5)
    run_ver.font.color.rgb = COLOR_MUTED

    doc.add_page_break()

    # Parse content
    i = 0
    in_code_block = False
    code_lines = []

    # Skip initial title lines from markdown since we added the cover
    while i < len(lines):
        line = lines[i]
        trimmed = line.strip()

        # Check code fence
        if trimmed.startswith('```'):
            if in_code_block:
                # Flush code block
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                p_code = cell.paragraphs[0]
                p_code.paragraph_format.space_before = Pt(0)
                p_code.paragraph_format.space_after = Pt(0)
                p_code.paragraph_format.line_spacing = 1.15
                code_text = "".join(code_lines).rstrip()
                r_code = p_code.add_run(code_text)
                r_code.font.name = 'Consolas'
                r_code.font.size = Pt(9.5)
                r_code.font.color.rgb = RGBColor(15, 23, 42)
                
                # Add spacing after table
                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_before = Pt(4)
                p_after.paragraph_format.space_after = Pt(6)

                in_code_block = False
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Skip main markdown H1/H2 cover headers
        if (trimmed.startswith('# Stackly Travel') or 
            trimmed.startswith('## Technical Documentation') or 
            trimmed.startswith('*Responsive HTML5 Template') or 
            trimmed.startswith('**Version 1.0**') or 
            trimmed == '---'):
            i += 1
            continue

        # Table detection
        if '|' in trimmed and ('|' in lines[i+1] if i+1 < len(lines) else False):
            # Gather table lines
            table_lines = []
            while i < len(lines) and '|' in lines[i].strip():
                table_lines.append(lines[i].strip())
                i += 1
            
            # Process table
            headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
            # skip separator line table_lines[1]
            rows_data = []
            for r_line in table_lines[2:]:
                cols = [c.strip() for c in r_line.split('|')[1:-1]]
                if any(cols):
                    rows_data.append(cols)
            
            if headers and rows_data:
                table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False

                # Format Header Row
                hdr_cells = table.rows[0].cells
                for idx, h_text in enumerate(headers):
                    cell = hdr_cells[idx]
                    set_cell_background(cell, "0B3B24")
                    set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    clean_h = re.sub(r'[*_`]', '', h_text)
                    r = p.add_run(clean_h)
                    r.font.name = 'Segoe UI Semibold'
                    r.font.bold = True
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor(255, 255, 255)

                # Format Body Rows
                for row_idx, r_cols in enumerate(rows_data):
                    row_cells = table.rows[row_idx + 1].cells
                    bg_color = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
                    for col_idx, c_text in enumerate(r_cols):
                        if col_idx < len(row_cells):
                            cell = row_cells[col_idx]
                            set_cell_background(cell, bg_color)
                            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(0)
                            p.paragraph_format.space_after = Pt(0)
                            clean_c = re.sub(r'[`]', '', c_text)
                            # Handle bold
                            parts = re.split(r'(\*\*.*?\*\*)', clean_c)
                            for part in parts:
                                if part.startswith('**') and part.endswith('**'):
                                    r = p.add_run(part[2:-2])
                                    r.font.bold = True
                                else:
                                    r = p.add_run(part)
                                r.font.name = 'Segoe UI'
                                r.font.size = Pt(9.5)
                                r.font.color.rgb = COLOR_DARK

                p_spc = doc.add_paragraph()
                p_spc.paragraph_format.space_before = Pt(4)
                p_spc.paragraph_format.space_after = Pt(6)
            continue

        # Headings
        if trimmed.startswith('## '):
            h_text = trimmed[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = 'Segoe UI Semibold'
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = COLOR_PRIMARY
            add_horizontal_border(p, color_hex="1B4D3E", sz="12")
            i += 1
            continue

        if trimmed.startswith('### '):
            h_text = trimmed[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = 'Segoe UI Semibold'
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        if trimmed.startswith('#### '):
            h_text = trimmed[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = 'Segoe UI Semibold'
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_SECONDARY
            i += 1
            continue

        # List items
        if trimmed.startswith('- ') or trimmed.startswith('* '):
            item_text = trimmed[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            # Parse formatting
            parts = re.split(r'(\*\*.*?\*\*|\`.*?\`)', item_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                elif part.startswith('`') and part.endswith('`'):
                    r = p.add_run(part[1:-1])
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(194, 65, 12)
                else:
                    r = p.add_run(part)
                r.font.name = 'Segoe UI'
                r.font.size = Pt(10.5)
            i += 1
            continue

        # Numbered list
        m_num = re.match(r'^(\d+)\.\s+(.*)$', trimmed)
        if m_num:
            num_str = m_num.group(1)
            item_text = m_num.group(2).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            parts = re.split(r'(\*\*.*?\*\*|\`.*?\`)', item_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                elif part.startswith('`') and part.endswith('`'):
                    r = p.add_run(part[1:-1])
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(194, 65, 12)
                else:
                    r = p.add_run(part)
                r.font.name = 'Segoe UI'
                r.font.size = Pt(10.5)
            i += 1
            continue

        # Empty line
        if not trimmed:
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.18
        parts = re.split(r'(\*\*.*?\*\*|\`.*?\`|\*.*?\*)', trimmed)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.font.bold = True
            elif part.startswith('*') and part.endswith('*') and len(part) > 2:
                r = p.add_run(part[1:-1])
                r.font.italic = True
            elif part.startswith('`') and part.endswith('`'):
                r = p.add_run(part[1:-1])
                r.font.name = 'Consolas'
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(194, 65, 12)
            else:
                r = p.add_run(part)
            r.font.name = 'Segoe UI'
            r.font.size = Pt(10.5)
            r.font.color.rgb = COLOR_DARK
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")

if __name__ == '__main__':
    md_file = r"c:\Users\prani\Desktop\Stackly Projects\Travel & Tourism\TECHNICAL_DOCUMENTATION.md"
    docx_file = r"c:\Users\prani\Desktop\Stackly Projects\Travel & Tourism\TECHNICAL_DOCUMENTATION.docx"
    build_docx(md_file, docx_file)

